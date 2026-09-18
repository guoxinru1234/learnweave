#!/usr/bin/env python3
"""Ingest course lectures and experiment manuals into a structured asset index.

The script is intentionally local and deterministic. It extracts text from
Markdown, DOCX, DOC, PDF, and TXT files, writes JSON/JSONL indexes, and can
optionally build a local ChromaDB index with a hashing embedding function.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
except Exception:  # pragma: no cover - reported during runtime
    Document = None  # type: ignore[assignment]


ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = ROOT.parent
KB_DIR = ROOT / "knowledge-base"
ASSET_ROOT = REPO_ROOT / "实验手册"
OUT_DIR = KB_DIR / "assets"
CHROMA_DIR = KB_DIR / ".chroma"

SUPPORTED_SUFFIXES = {".docx", ".doc", ".pdf", ".txt", ".md"}

TOPIC_RULES = [
    ("scala", "Scala"),
    ("函数式", "函数式编程"),
    ("编程实战", "编程实践"),
    ("spark sql", "Spark SQL"),
    ("sparksql", "Spark SQL"),
    ("dataframe", "DataFrame"),
    ("rdd", "RDD"),
    ("yarn", "YARN"),
    ("standalone", "Standalone"),
    ("local", "Local"),
    ("流处理", "流处理"),
    ("stream", "流处理"),
    ("开发环境", "环境部署"),
    ("环境部署", "环境部署"),
    ("部署", "环境部署"),
    ("读写数据", "数据读写"),
    ("作业", "作业"),
    ("习题", "习题"),
    ("答案", "答案"),
]


def repo_rel(path: Path) -> str:
    return path.resolve().relative_to(REPO_ROOT).as_posix()


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    compact: list[str] = []
    blank = False
    for line in lines:
        if not line:
            if not blank:
                compact.append("")
            blank = True
            continue
        compact.append(line)
        blank = False
    return "\n".join(compact).strip()


def extract_docx(path: Path) -> str:
    if Document is None:
        raise RuntimeError("python-docx is required to extract DOCX files")
    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return clean_text("\n".join(parts))


def extract_doc(path: Path) -> str:
    if not shutil.which("textutil"):
        raise RuntimeError("textutil is required to extract legacy DOC files on macOS")
    with tempfile.TemporaryDirectory() as tmp:
        out_path = Path(tmp) / f"{path.stem}.txt"
        subprocess.run(
            ["textutil", "-convert", "txt", "-output", str(out_path), str(path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        return clean_text(out_path.read_text(encoding="utf-8", errors="ignore"))


def extract_pdf(path: Path) -> str:
    if shutil.which("pdftotext"):
        result = subprocess.run(
            ["pdftotext", "-layout", str(path), "-"],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        return clean_text(result.stdout.decode("utf-8", errors="ignore"))
    try:
        from pypdf import PdfReader
    except Exception as exc:  # pragma: no cover - environment fallback
        raise RuntimeError("pdftotext or pypdf is required to extract PDF files") from exc
    reader = PdfReader(str(path))
    return clean_text("\n".join(page.extract_text() or "" for page in reader.pages))


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".doc":
        return extract_doc(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in {".txt", ".md"}:
        return clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    raise ValueError(f"Unsupported file type: {path}")


def classify(path: Path) -> str:
    name = path.stem.lower()
    rel = repo_rel(path)
    if path.suffix.lower() == ".txt":
        return "dataset"
    if "问题/" in rel or "/问题/" in rel:
        return "faq"
    if "案例/" in rel or "/案例/" in rel:
        return "case"
    if "答案" in name:
        return "answer"
    if "习题" in name or "作业" in name:
        return "exercise"
    if any(key in name for key in ["实战", "部署", "环境", "创建", "读写", "流处理"]):
        return "lab"
    return "manual"


def detect_topics(path: Path, text: str) -> list[str]:
    haystack = f"{path.stem} {text[:2000]}".lower()
    topics = [label for key, label in TOPIC_RULES if key in haystack]
    seen: set[str] = set()
    return [topic for topic in topics if not (topic in seen or seen.add(topic))]


def infer_lecture_ids(topics: Iterable[str]) -> list[int]:
    topic_set = set(topics)
    ids: set[int] = set()
    if {"Scala", "函数式编程", "编程实践", "环境部署"} & topic_set:
        ids.update(range(1, 7))
    if {"Local", "Standalone", "YARN", "环境部署"} & topic_set:
        ids.update(range(7, 11))
    if "RDD" in topic_set:
        ids.update(range(11, 17))
    if {"Spark SQL", "DataFrame", "数据读写"} & topic_set:
        ids.update(range(17, 21))
    return sorted(ids)


def chunk_text(text: str, max_chars: int = 1200, overlap: int = 160) -> list[str]:
    paragraphs = [p.strip() for p in re.split(r"\n{2,}", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(current) + len(paragraph) + 2 <= max_chars:
            current = f"{current}\n\n{paragraph}".strip()
            continue
        if current:
            chunks.append(current)
        while len(paragraph) > max_chars:
            chunks.append(paragraph[:max_chars])
            paragraph = paragraph[max_chars - overlap :]
        current = paragraph
    if current:
        chunks.append(current)
    return chunks


def stable_asset_id(path: Path, index: int) -> str:
    digest = hashlib.sha1(repo_rel(path).encode("utf-8")).hexdigest()[:8]
    return f"asset-{index:03d}-{digest}"


def load_lectures() -> tuple[list[dict], list[dict]]:
    index_path = KB_DIR / "index.json"
    index_data = json.loads(index_path.read_text(encoding="utf-8"))
    lecture_assets: list[dict] = []
    chunks: list[dict] = []
    for lecture in index_data.get("lectures", []):
        lecture_path = KB_DIR / lecture["dir"] / lecture["file"]
        text = extract_text(lecture_path)
        asset_id = f"lecture-{int(lecture['id']):02d}"
        asset = {
            "id": asset_id,
            "title": lecture["title"].strip(),
            "category": "lecture",
            "source_path": repo_rel(lecture_path),
            "suffix": ".md",
            "topics": detect_topics(lecture_path, text),
            "lecture_ids": [int(lecture["id"])],
            "text_chars": len(text),
            "chunk_count": 0,
        }
        lecture_chunks = chunk_text(text)
        asset["chunk_count"] = len(lecture_chunks)
        lecture_assets.append(asset)
        for chunk_index, chunk in enumerate(lecture_chunks):
            chunks.append(chunk_record(asset, chunk_index, chunk))
    return lecture_assets, chunks


def chunk_record(asset: dict, chunk_index: int, text: str) -> dict:
    return {
        "id": f"{asset['id']}:{chunk_index:03d}",
        "asset_id": asset["id"],
        "title": asset["title"],
        "category": asset["category"],
        "source_path": asset["source_path"],
        "topics": asset["topics"],
        "lecture_ids": asset["lecture_ids"],
        "chunk_index": chunk_index,
        "text": text,
    }


def ingest_assets() -> tuple[dict, list[dict], list[dict]]:
    assets: list[dict] = []
    chunks: list[dict] = []
    files = sorted(
        path
        for path in ASSET_ROOT.rglob("*")
        if path.is_file()
        and path.name != ".DS_Store"
        and path.suffix.lower() in SUPPORTED_SUFFIXES
    )
    for index, path in enumerate(files, start=1):
        try:
            text = extract_text(path)
        except Exception as e:
            print(f"[WARN] 跳过无法提取的文件 {path.name}: {e}")
            continue
        category = classify(path)
        topics = detect_topics(path, text)
        asset = {
            "id": stable_asset_id(path, index),
            "title": path.stem.strip(),
            "category": category,
            "source_path": repo_rel(path),
            "suffix": path.suffix.lower(),
            "topics": topics,
            "lecture_ids": infer_lecture_ids(topics),
            "text_chars": len(text),
            "chunk_count": 0,
        }
        asset_chunks = chunk_text(text)
        asset["chunk_count"] = len(asset_chunks)
        assets.append(asset)
        for chunk_index, chunk in enumerate(asset_chunks):
            chunks.append(chunk_record(asset, chunk_index, chunk))
    lecture_assets, lecture_chunks = load_lectures()
    all_assets = lecture_assets + assets
    all_chunks = lecture_chunks + chunks
    catalog = {
        "schema_version": 1,
        "generated_by": "scripts/ingest_course_assets.py",
        "course": "大数据计算集群技术",
        "asset_root": repo_rel(ASSET_ROOT),
        "counts": {
            "assets": len(assets),
            "lectures": len(lecture_assets),
            "chunks": len(all_chunks),
            "by_category": dict(sorted(Counter(asset["category"] for asset in all_assets).items())),
        },
        "assets": all_assets,
    }
    return catalog, assets, all_chunks


def write_outputs(catalog: dict, experiment_assets: list[dict], chunks: list[dict]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "catalog.json").write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    with (OUT_DIR / "chunks.jsonl").open("w", encoding="utf-8") as f:
        for chunk in chunks:
            f.write(json.dumps(chunk, ensure_ascii=False) + "\n")
    labs = [asset for asset in experiment_assets if asset["category"] in {"lab", "case"}]
    datasets = [asset for asset in experiment_assets if asset["category"] == "dataset"]
    (OUT_DIR / "experiments.json").write_text(
        json.dumps({"total": len(labs), "experiments": labs}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (OUT_DIR / "datasets.json").write_text(
        json.dumps({"total": len(datasets), "datasets": datasets}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (OUT_DIR / "README.md").write_text(
        "# 知识资产索引\n\n"
        "本目录由 `scripts/ingest_course_assets.py` 生成，包含课程讲义、实验手册、作业、案例、FAQ 和数据集的结构化索引。\n\n"
        "- `catalog.json`: 全量资产元数据\n"
        "- `chunks.jsonl`: RAG 检索文本块\n"
        "- `experiments.json`: 实验与案例资产\n"
        "- `datasets.json`: 数据集资产\n\n"
        "本地 ChromaDB 索引可通过 `make assets-chroma` 生成到 `knowledge-base/.chroma/`，该目录不提交 Git。\n",
        encoding="utf-8",
    )


class HashEmbeddingFunction:
    """Small deterministic embedding function for local ChromaDB indexes."""

    def __init__(self, dimensions: int = 384):
        self.dimensions = dimensions

    def name(self) -> str:
        return f"learnmate-hash-{self.dimensions}"

    def __call__(self, input):  # Chroma validates this parameter name.
        return [self.embed(text) for text in input]

    def embed(self, text: str) -> list[float]:
        vector = [0.0] * self.dimensions
        tokens = re.findall(r"[A-Za-z0-9_]+|[\u4e00-\u9fff]", text.lower())
        for token in tokens:
            digest = hashlib.sha1(token.encode("utf-8")).digest()
            bucket = int.from_bytes(digest[:4], "big") % self.dimensions
            sign = 1.0 if digest[4] % 2 == 0 else -1.0
            vector[bucket] += sign
        norm = math.sqrt(sum(value * value for value in vector)) or 1.0
        return [value / norm for value in vector]


def build_chroma(chunks: list[dict]) -> None:
    try:
        import chromadb
    except Exception as exc:
        raise RuntimeError("chromadb is required for --build-chroma") from exc
    if CHROMA_DIR.exists():
        shutil.rmtree(CHROMA_DIR)
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    collection = client.get_or_create_collection(
        "learnmate_course_assets",
        embedding_function=HashEmbeddingFunction(),
        metadata={"hnsw:space": "cosine"},
    )
    batch_size = 256
    for start in range(0, len(chunks), batch_size):
        batch = chunks[start : start + batch_size]
        collection.add(
            ids=[item["id"] for item in batch],
            documents=[item["text"] for item in batch],
            metadatas=[
                {
                    "asset_id": item["asset_id"],
                    "title": item["title"],
                    "category": item["category"],
                    "source_path": item["source_path"],
                    "topics": ",".join(item["topics"]),
                    "lecture_ids": ",".join(str(x) for x in item["lecture_ids"]),
                }
                for item in batch
            ],
        )
    (OUT_DIR / "chroma_manifest.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "generated_by": "scripts/ingest_course_assets.py --build-chroma",
                "path": str(CHROMA_DIR.relative_to(ROOT)),
                "collection": "learnmate_course_assets",
                "embedding": "learnmate-hash-384",
                "chunks": len(chunks),
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--build-chroma", action="store_true", help="build local ChromaDB index")
    args = parser.parse_args()

    catalog, experiment_assets, chunks = ingest_assets()
    write_outputs(catalog, experiment_assets, chunks)
    if args.build_chroma:
        build_chroma(chunks)
    print(
        json.dumps(
            {
                "assets": catalog["counts"]["assets"],
                "lectures": catalog["counts"]["lectures"],
                "chunks": catalog["counts"]["chunks"],
                "categories": catalog["counts"]["by_category"],
                "chroma": args.build_chroma,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
